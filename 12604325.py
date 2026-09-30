import openpyxl
import datetime
import math

# Loads daily activity data from the Excel sheet into a list of dictionaries
def load_data(filename):
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
    except FileNotFoundError:
        print(f"Error: Could not find the file '{filename}'. Check the filename and that it's in the right folder.")
        return []

    ws = wb["Daily Log"]

    all_days = []
    for row in range(6, ws.max_row + 1):
        day = {
            "date": ws.cell(row, 1).value,
            "sleep": ws.cell(row, 2).value,
            "fitness": ws.cell(row, 3).value,
            "study": ws.cell(row, 4).value,
            "coding": ws.cell(row, 5).value,
            "class_time": ws.cell(row, 6).value,
            "classes_attended": ws.cell(row, 7).value,
            "other": ws.cell(row, 8).value,
            "total_tracked": ws.cell(row, 9).value,
            "free_time": ws.cell(row, 10).value,
            "feeling": ws.cell(row, 11).value,
            "satisfaction": ws.cell(row, 12).value,
            "energy": ws.cell(row, 13).value,
        }
        all_days.append(day)

    return all_days

all_days = load_data("12604325.xlsx")
print(len(all_days))
print(all_days[0])

cutoff_date = datetime.datetime(2026, 9, 21)

# Checks whether one day's data is complete and within the recording period
def check_valid(day):
    if day["date"] is None:
        return False
    if day["date"] > cutoff_date:
        return False
    if day["sleep"] is None:
        return False
    if day["fitness"] is None:
        return False
    if day["study"] is None:
        return False
    if day["coding"] is None:
        return False
    if day["class_time"] is None:
        return False
    if day["other"] is None:
        return False
    if day["total_tracked"] is None:
        return False
    if day["free_time"] is None:
        return False
    if day["feeling"] is None:
        return False
    if day["satisfaction"] is None:
        return False
    if day["energy"] is None:
        return False
    return True
 
# Counts how many days are valid vs missing, compared to expected days
def count_days(all_days, expected_days):
    valid_count = 0
    missing_count = 0

    for day in all_days:
        if check_valid(day):
            valid_count += 1
        else:
            missing_count += 1

    return valid_count, missing_count, expected_days

valid, missing, expected = count_days(all_days, 40)
print("Valid:", valid)
print("Missing:", missing)
print("Expected:", expected)

feeling_scale = {
    "Stressed": 1,
    "Low": 2,
    "Neutral": 3,
    "Good": 4,
    "Excellent": 5,
}

satisfaction_scale = {
    "Unsatisfied": 1,
    "Neutral": 2,
    "Satisfied": 3,
    "Very Satisfied": 4,
}

energy_scale = {
    "Very Low": 1,
    "Low": 2,
    "Medium": 3,
    "High": 4,
    "Very High": 5,
}

# Calculates average minutes for each activity, using only valid days
def calculate_averages(all_days):
    total_sleep = 0
    total_fitness = 0
    total_study = 0
    total_coding = 0
    total_class = 0
    total_other = 0
    total_free = 0
    total_tracked = 0
    total_ei_sum = 0
    valid_count = 0

    for day in all_days:
        if check_valid(day):
            total_sleep += day["sleep"]
            total_fitness += day["fitness"]
            total_study += day["study"]
            total_coding += day["coding"]
            total_class += day["class_time"]
            total_other += day["other"]
            total_free += day["free_time"]
            total_tracked += day["total_tracked"]
            
            feeling_score = feeling_scale[day["feeling"]]
            satisfaction_score = satisfaction_scale[day["satisfaction"]]
            energy_score = energy_scale[day["energy"]]
            total_ei_sum += (feeling_score + satisfaction_score + energy_score)
            
            valid_count += 1

    averages = {
        "sleep": total_sleep / valid_count,
        "fitness": total_fitness / valid_count,
        "study": total_study / valid_count,
        "coding": total_coding / valid_count,
        "class_time": total_class / valid_count,
        "other": total_other / valid_count,
        "free_time": total_free / valid_count,
        "total_tracked": total_tracked / valid_count,
        "ei": total_ei_sum / (3 * valid_count),
    }

    return averages

averages = calculate_averages(all_days)
print(averages)

# Plugs averages into the 9 formulas (TPI, AAI, PhAI, SRI, ABI, TUI, EI, DCI, PAI)
def calculate_indices(averages, valid_days, expected_days):
    tpi = averages["coding"]
    aai = averages["study"] + averages["class_time"]
    phai = averages["fitness"]
    sri = averages["sleep"]
    abi = averages["free_time"]
    tui = averages["total_tracked"]
    ei = averages["ei"]
    dci = (valid_days / expected_days) * 100

    pai = (0.15 * tpi) + (0.20 * aai) + (0.15 * phai) + (0.20 * sri) + (0.15 * tui) + (0.10 * ei) + (0.05 * dci)

    indices = {
        "TPI": tpi,
        "AAI": aai,
        "PhAI": phai,
        "SRI": sri,
        "ABI": abi,
        "TUI": tui,
        "EI": ei,
        "DCI": dci,
        "PAI": pai,
    }

    return indices

indices = calculate_indices(averages, valid, expected)
print(indices)

# Calculates Pearson correlation between two numeric lists, without NumPy
def calculate_correlation(x_values, y_values):
    n = len(x_values)
    sum_x = sum(x_values)
    sum_y = sum(y_values)
    sum_xy = 0
    sum_x2 = 0
    sum_y2 = 0
    
    for i in range(n):
        sum_xy += x_values[i] * y_values[i]
        sum_x2 += x_values[i] ** 2
        sum_y2 += y_values[i] ** 2
    
    numerator = (n * sum_xy) - (sum_x * sum_y)
    denominator = math.sqrt((n * sum_x2 - sum_x ** 2) * (n * sum_y2 - sum_y ** 2))
    
    if denominator == 0:
        return 0
    
    return numerator / denominator

# Pulls matching value pairs from valid days, converting words to numbers if needed
def get_pairs(all_days, field1, field2, scale2=None):
    x_values = []
    y_values = []
    
    for day in all_days:
        if check_valid(day):
            x = day[field1]
            y = day[field2]
            if scale2 is not None:
                y = scale2[y]
            x_values.append(x)
            y_values.append(y)
    
    return x_values, y_values


coding_vals, energy_vals = get_pairs(all_days, "coding", "energy", energy_scale)
sleep_vals, energy_vals2 = get_pairs(all_days, "sleep", "energy", energy_scale)
study_vals, satisfaction_vals = get_pairs(all_days, "study", "satisfaction", satisfaction_scale)

coding_energy_corr = calculate_correlation(coding_vals, energy_vals)
sleep_energy_corr = calculate_correlation(sleep_vals, energy_vals2)
study_satisfaction_corr = calculate_correlation(study_vals, satisfaction_vals)

print("Coding <-> Energy correlation:", coding_energy_corr)
print("Sleep <-> Energy correlation:", sleep_energy_corr)
print("Study <-> Satisfaction correlation:", study_satisfaction_corr)

print("\n--- FINAL SUMMARY ---")
print(f"Valid days: {valid} / Expected: {expected} (Missing: {missing})")
print(f"TPI: {indices['TPI']:.2f}")
print(f"AAI: {indices['AAI']:.2f}")
print(f"PhAI: {indices['PhAI']:.2f}")
print(f"SRI: {indices['SRI']:.2f}")
print(f"ABI: {indices['ABI']:.2f}")
print(f"TUI: {indices['TUI']:.2f}")
print(f"EI: {indices['EI']:.2f}")
print(f"DCI: {indices['DCI']:.2f}")
print(f"PAI: {indices['PAI']:.2f}")